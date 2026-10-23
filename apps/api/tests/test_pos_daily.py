"""POS 每日營業報表：營收入損益（內用 R01、外帶 R02，未稅），每種收款方式一張應收，銀行入帳對上。"""
import io
import uuid

from openpyxl import Workbook

from tests.test_bank_webhook import HOOK, _send, _txn

TENDERS = ["現金", "信用卡(Global Payment TW)", "LINE Pay(LINE Pay)", "禮券", "台灣pay", "手刷信用卡"]


def _sheet(ws, extra_cols, rows):
    """rows: [(日期, {收款方式: 金額})]，格式照 POS 下載的檔案（收款方式在「日期」上一列，金額在 NT$ 欄）。"""
    head = ["日期", *extra_cols]
    above = [None] * len(head)
    for t in TENDERS:
        above += [t, None]
        head += ["數量", "NT$"]
    head.append("總收入")
    ws.append(above)
    ws.append(head)
    for d, amounts in rows:
        r = [d] + [0] * len(extra_cols)
        for t in TENDERS:
            r += [1 if amounts.get(t) else 0, amounts.get(t, 0)]
        r.append(sum(amounts.values()))
        ws.append(r)


def pos_file(days):
    wb = Workbook()
    s = wb.active
    s.title = "Sales Summary"
    s.append(["2026-10-01 00:00 - 2026-10-02 23:59"])
    s.append(["時段: 00:00:00 - 23:59:59"])
    _sheet(s, ["銷售額", "服務費", "送遞附加費", "折扣額", "退款", "總銷售額", "預訂單訂金", "預定單訂金退回", "禮券退款",
               "進位差異", "小費", "客戶", "平均消費", "開枱率"],
           [(d, {t: a + b for t, (a, b) in v.items()}) for d, v in days])
    _sheet(wb.create_sheet("內用"), ["銷售額", "服務費", "折扣額", "退款", "總銷售額", "進位差異", "小費", "客戶", "平均消費", "開枱率"],
           [(d, {t: a for t, (a, b) in v.items()}) for d, v in days])
    _sheet(wb.create_sheet("外帶"), ["銷售額", "送遞附加費", "折扣額", "退款", "總銷售額", "進位差異", "小費", "客戶"],
           [(d, {t: b for t, (a, b) in v.items()}) for d, v in days])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DAYS = [("2026-10-01 (四)", {"現金": (19855, 35105), "信用卡(Global Payment TW)": (7265, 17619), "LINE Pay(LINE Pay)": (1525, 8097)}),
        ("2026-10-02 (五)", {"現金": (13515, 29330), "信用卡(Global Payment TW)": (2655, 1348)})]


def _upload(client, data, store="store-ys-01"):
    r = client.post("/api/v1/imports", data={"kind": "pos_daily", "company_id": "comp-2", "period": "2026-10", "store_id": store},
                    files={"file": ("營業報表.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 201, r.text
    return r.json()


def test_pos_import_revenue_and_cash_deposit(client, svc_headers):
    b = _upload(client, pos_file(DAYS))
    assert b["kind"] == "pos_daily" and b["row_count"] == 5 and not b["errors"]
    assert b["total_amount"] == 54960 + 24884 + 9622 + 42845 + 4003
    r = client.post(f"/api/v1/imports/{b['id']}/commit")
    assert r.status_code == 200 and r.json()["created"] == 5, r.text
    pl = client.get("/api/v1/pl/statement", params={"period_from": "2026-10", "period_to": "2026-10", "company_id": "comp-2",
                                                    "store_id": "store-ys-01", "columns": "month"}).json()
    rev = {l["category_code"]: l["values"]["total"] for s in pl["sections"] for l in s["lines"]}
    assert rev["R01"] == round(19855 / 1.05) + round(7265 / 1.05) + round(1525 / 1.05) + round(13515 / 1.05) + round(2655 / 1.05)
    assert "R02" in rev
    # 重新匯入同一份：全部略過
    b2 = _upload(client, pos_file(DAYS))
    assert {p["action"] for p in b2["preview"]} == {"skip"}
    assert client.post(f"/api/v1/imports/{b2['id']}/commit").json()["skipped"] == 5
    # 10/2 現金存進銀行 54,960 → 自動對上 10/1 的現金
    res = _send(client, svc_headers, _txn("DEP1", 54960, direction="in", cash_account_id="acc-3312", txn_date="2026-10-02",
                                          bank_memo="現金存入"))
    assert res["matched"] == 1, res
    # 信用卡撥款（兩天合併、扣 2% 手續費）→ 待確認 → 手動核對、差額記刷卡手續費
    card_total = 24884 + 4003
    net = card_total - round(card_total * 0.02)
    res = _send(client, svc_headers, _txn("CARD1", net, direction="in", cash_account_id="acc-3312", txn_date="2026-10-04",
                                          bank_memo="Global Payment"))
    assert res["review"] == 1
    docs = client.get("/api/v1/documents", params={"doc_type": "receivable", "company_id": "comp-2", "q": "Global"}).json()["items"]
    item = next(i for i in client.get("/api/v1/review-items", params={"kind": "bank_only"}).json()["items"]
                if i["payload"]["item"]["cash_txn_id"] == "CARD1")
    r = client.post(f"/api/v1/review-items/{item['id']}/resolve",
                    json={"action": "manual_match", "payload": {"document_ids": [d["id"] for d in docs], "difference_mode": "book"}})
    assert r.status_code == 200, r.text
    fee = client.get(f"/api/v1/documents/{r.json()['difference_document']}").json()
    assert fee["lines"][0]["category_code"] == "E11" and fee["settle_amount"] == round(card_total * 0.02)


def test_pos_requires_store_and_auto_detects(client):
    r = client.post("/api/v1/imports", data={"kind": "pos_daily", "company_id": "comp-2", "period": "2026-10"},
                    files={"file": ("x.xlsx", pos_file(DAYS), "application/octet-stream")})
    assert r.status_code == 422 and "門市" in r.text
    # 選錯成「應付表」也會自動認出是 POS 報表
    r = client.post("/api/v1/imports", data={"kind": "ap_sheet", "company_id": "comp-2", "period": "2026-10", "store_id": "store-ys-01"},
                    files={"file": ("x.xlsx", pos_file(DAYS), "application/octet-stream")})
    assert r.status_code == 201 and r.json()["kind"] == "pos_daily", r.text
