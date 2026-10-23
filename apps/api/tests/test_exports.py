"""Excel 下載：損益表、應付／應收／預收。"""
import io

from openpyxl import load_workbook

from tests.helpers import payable
from tests.test_pl import _receivable


def _wb(r):
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    return load_workbook(io.BytesIO(r.content))


def test_pl_export_matches_statement(client):
    _receivable(client, "comp-2", "store-ys-01", "2026-08", 100_000)
    payable(client, company="comp-2", store="store-ys-01", period="2026-08", amount=31_500, category="C01")
    params = {"period_from": "2026-08", "period_to": "2026-08", "company_id": "comp-2", "columns": "store"}
    st = client.get("/api/v1/pl/statement", params=params).json()
    wb = _wb(client.get("/api/v1/pl/statement/export", params=params))
    ws = wb["損益表"]
    rows = {r[0]: r for r in ws.iter_rows(values_only=True) if r and r[0]}
    assert "楊師傅滷味有限公司 損益表" in rows
    assert rows["毛利"][-2] == st["subtotals"]["gross_profit"]["total"] == 100_000 - 30_000
    assert "明細" in wb.sheetnames and wb["明細"].max_row >= 6
    # 依月份
    wb = _wb(client.get("/api/v1/pl/statement/export", params={**params, "period_from": "2026-07", "columns": "month"}))
    assert "增減" in [c.value for c in wb["損益表"][4]]


def test_documents_export(client):
    payable(client, company="comp-2", store="store-ys-01", period="2026-08", amount=10_500, vendor="元華行", invoice_no="EX00000001")
    payable(client, company="comp-2", store=None, period="2026-08", amount=2_100, category="E09", vendor="富霖印刷", invoice_no="EX00000002")
    r = client.get("/api/v1/documents/export", params={"doc_type": "payable", "company_id": "comp-2", "period": "2026-08"})
    wb = _wb(r)
    ws = wb["應付"]
    data = [row for row in ws.iter_rows(min_row=5, values_only=True)]
    assert len(data) == 3 and data[-1][0] == "合計" and str(data[-1][10]).startswith("=SUM(")
    assert {row[4] for row in data[:2]} == {"元華行", "富霖印刷"}
    assert "公司層級" in {row[3] for row in data[:2]}
    assert wb["明細"].max_row == 6
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
