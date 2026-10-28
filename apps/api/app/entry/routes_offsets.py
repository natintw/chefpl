"""v0.6 使用者端 API：廠商主檔、預付款、折讓、代付、沖抵（開發文件 5.13–5.17）。

路由只做：取得使用者 → 權限 → 呼叫功能層 → 轉成回應。
"""
from __future__ import annotations

import io
import json
import uuid
from datetime import date, timedelta

from fastapi import APIRouter, Depends, File, Form, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.db import get_ledger_session
from app.entry import serializers as ser
from app.entry.deps import need
from app.entry.routes import _check_company, _cids, _doc_or_404
from app.entry.schemas import AllowanceIn, OffsetIn, OnBehalfIn, PrepaymentIn, ReasonIn, VendorIn, VendorPatch
from app.errors import NotFound
from app.feature.offsets import OffsetService
from app.feature.posting import PostingService
from app.feature.repositories import DocumentRepository
from app.feature.vendors import VendorMasterService
from app.guard.principals import UserPrincipal

router = APIRouter(prefix="/api/v1")
SESSION = Depends(get_ledger_session, scope="function")


def _vendor(c) -> dict:
    return {"id": str(c.id), "name": c.name, "tax_id": c.tax_id, "no_tax_id": c.no_tax_id,
            "id_masked": c.id_masked, "aliases": c.aliases or [], "keywords": c.keywords or [],
            "default_category_code": c.default_category_code, "is_intercompany": c.is_intercompany,
            "is_active": c.is_active}


# ---------------------------------------------------------------- 廠商主檔（S13）
@router.get("/vendors/search")
def vendor_search(q: str = Query(min_length=1), company_id: str | None = None, limit: int = Query(10, le=50),
                  user: UserPrincipal = Depends(need("doc.view")), s: Session = SESSION):
    """統編（8 碼數字或開頭）或名稱、簡稱、付款關鍵字；附未沖抵預付金額（5.16）。"""
    return {"items": VendorMasterService(s).search(q, limit=limit, company_id=company_id)}


@router.get("/vendors")
def vendor_list(q: str | None = None, missing_tax_id: bool = False, include_inactive: bool = False,
                user: UserPrincipal = Depends(need("doc.view")), s: Session = SESSION):
    items = VendorMasterService(s).list(q, missing_tax_id, include_inactive)
    return {"items": items, "missing_tax_id_count": sum(1 for i in items if i["missing_tax_id"])}


@router.post("/vendors", status_code=201)
def vendor_create(body: VendorIn, user: UserPrincipal = Depends(need("settings.manage")), s: Session = SESSION):
    return _vendor(VendorMasterService(s).create(**body.model_dump()))


@router.patch("/vendors/{cp_id}")
def vendor_patch(cp_id: uuid.UUID, body: VendorPatch, user: UserPrincipal = Depends(need("settings.manage")),
                 s: Session = SESSION):
    return _vendor(VendorMasterService(s).update(cp_id, **body.model_dump(exclude_unset=True)))


@router.post("/vendors/import")
async def vendor_import(file: UploadFile = File(...), commit: bool = Form(False), decisions: str = Form("{}"),
                        user: UserPrincipal = Depends(need("settings.manage")), s: Session = SESSION):
    """commit=false 只預覽；commit=true 寫入。decisions：{"列號": "update_name"|"keep_name"|"skip"}。"""
    dec = {int(k): v for k, v in (json.loads(decisions or "{}") or {}).items()}
    return VendorMasterService(s).import_xlsx(await file.read(), commit=commit, decisions=dec).to_dict()


@router.get("/vendors/import-template")
def vendor_template(user: UserPrincipal = Depends(need("doc.view"))):
    wb = Workbook()
    ws = wb.active
    ws.title = "匯入名單"
    ws.append(["廠商名稱", "廠商統編", "沒有統編", "遮蔽身分證", "簡稱", "預設類別", "付款關鍵字", "狀態"])
    ws.append(["富統食品股份有限公司", "23344294", "", "", "富統", "C02", "富統", "啟用"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E8EEF6")
    for row in ws.iter_rows(min_row=2, max_col=2):
        row[1].number_format = "@"
    for letter, w in zip("ABCDEFGH", (34, 12, 9, 14, 16, 10, 14, 8)):
        ws.column_dimensions[letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''vendor-master-template.xlsx"})


# ---------------------------------------------------------------- 預付款（5.13）
@router.post("/prepayments", status_code=201)
def prepayment_create(body: PrepaymentIn, user: UserPrincipal = Depends(need("doc.create")), s: Session = SESSION):
    _check_company(user, body.company_id)
    d = OffsetService(s).create_prepayment(actor=user.actor, **body.model_dump())
    return ser.document(d)


@router.get("/prepayments")
def prepayment_list(company_id: str | None = None, counterparty_id: uuid.UUID | None = None,
                    status: str = "open", user: UserPrincipal = Depends(need("doc.view")), s: Session = SESSION):
    svc = OffsetService(s)
    cids = _cids(user)
    docs, _ = DocumentRepository(s).list(doc_type="prepayment", company_ids=cids, page_size=500)
    out = []
    for d in docs:
        if company_id and d.company_id != company_id:
            continue
        if counterparty_id and d.counterparty_id != counterparty_id:
            continue
        rem = svc.remaining(d)
        if status == "open" and rem <= 0:
            continue
        out.append({**ser.document(d, detail=False), "remaining": rem,
                    "tax_id": d.counterparty.tax_id if d.counterparty else None})
    return {"items": out, "total_remaining": sum(i["remaining"] for i in out)}


@router.get("/vendors/{cp_id}/open-prepayments")
def vendor_open_prepayments(cp_id: uuid.UUID, company_id: str, before: date | None = None,
                            user: UserPrincipal = Depends(need("doc.view")), s: Session = SESSION):
    _check_company(user, company_id)
    return {"items": OffsetService(s).open_prepayments(company_id, cp_id, before)}


# ---------------------------------------------------------------- 折讓（5.14）
@router.get("/documents/lookup")
def document_lookup(invoice_no: str = Query(min_length=4), seller_tax_id: str | None = None,
                    user: UserPrincipal = Depends(need("doc.view")), s: Session = SESSION):
    docs = OffsetService(s).lookup_invoice(invoice_no, seller_tax_id, _cids(user))
    return {"items": [_original(s, d) for d in docs]}


@router.get("/documents/original-candidates")
def original_candidates(counterparty_id: uuid.UUID, company_id: str | None = None, min_amount: int = 0,
                        months: int = Query(3, ge=1, le=24), user: UserPrincipal = Depends(need("doc.view")),
                        s: Session = SESSION):
    if company_id:
        _check_company(user, company_id)
    since = date.today() - timedelta(days=31 * months)
    docs = OffsetService(s).original_candidates(counterparty_id, company_id, min_amount, since)
    cids = _cids(user)
    return {"items": [_original(s, d) for d in docs if cids is None or d.company_id in cids]}


def _original(s: Session, d) -> dict:
    svc = OffsetService(s)
    return {**ser.document(d), "allowance_room": svc.allowance_room(d),
            "tax_id": d.counterparty.tax_id if d.counterparty else None}


@router.post("/allowances", status_code=201)
def allowance_create(body: AllowanceIn, user: UserPrincipal = Depends(need("doc.create")), s: Session = SESSION):
    if body.original_document_id:
        _doc_or_404(s, user, body.original_document_id)
    elif body.company_id:
        _check_company(user, body.company_id)
    data = body.model_dump()
    use = data.pop("period")
    period = f"{body.doc_date.year:04d}-{body.doc_date.month:02d}" if use == "current" else None
    d = OffsetService(s).create_allowance(actor=user.actor, period=period, **data)
    return ser.document(d)


# ---------------------------------------------------------------- 代付（5.15）
@router.post("/on-behalf", status_code=201)
def on_behalf_create(body: OnBehalfIn, user: UserPrincipal = Depends(need("doc.create")), s: Session = SESSION):
    _check_company(user, body.beneficiary_company_id)   # 兩家公司都要有權限
    _check_company(user, body.payer_company_id)
    return OffsetService(s).create_on_behalf(actor=user.actor, **body.model_dump())


@router.post("/on-behalf/{doc_id}/void")
def on_behalf_void(doc_id: uuid.UUID, body: ReasonIn, user: UserPrincipal = Depends(need("doc.void")),
                   s: Session = SESSION):
    _doc_or_404(s, user, doc_id)
    OffsetService(s).void_on_behalf(doc_id, body.reason, user.actor)
    return {"voided": True}


# ---------------------------------------------------------------- 沖抵與單據關聯
@router.post("/offsets", status_code=201)
def offset_create(body: OffsetIn, user: UserPrincipal = Depends(need("doc.edit")), s: Session = SESSION):
    _doc_or_404(s, user, body.to_document_id)
    _doc_or_404(s, user, body.from_document_id)
    off = OffsetService(s).apply(body.kind, body.from_document_id, body.to_document_id, body.amount, user.actor)
    return {"id": str(off.id), "amount": off.amount}


@router.post("/offsets/{offset_id}/reverse")
def offset_reverse(offset_id: uuid.UUID, body: ReasonIn, user: UserPrincipal = Depends(need("doc.edit")),
                   s: Session = SESSION):
    off = OffsetService(s).reverse(offset_id, body.reason, user.actor)
    return {"id": str(off.id), "reversed": True}


@router.get("/documents/{doc_id}/chain")
def document_chain(doc_id: uuid.UUID, user: UserPrincipal = Depends(need("doc.view")), s: Session = SESSION):
    """單據關聯（6.4）：同一筆交易的預付、發票、折讓、付款、代付另一半，依日期排列。"""
    root = _doc_or_404(s, user, doc_id)
    repo, svc = DocumentRepository(s), OffsetService(s)
    seen: dict = {}
    todo = [root]
    while todo and len(seen) < 30:
        d = todo.pop()
        if d is None or d.id in seen:
            continue
        seen[d.id] = d
        if d.related_document_id:
            todo.append(repo.get(d.related_document_id))
        todo.extend(repo.related(d.id))
        for o in svc.list_for(d.id):
            if o["reversed_at"]:
                continue
            todo.append(repo.get(uuid.UUID(o["from_document_id"])))
            todo.append(repo.get(uuid.UUID(o["to_document_id"])))
    cids = _cids(user)
    docs = sorted((d for d in seen.values() if d.status != "void" and (cids is None or d.company_id in cids)),
                  key=lambda d: (d.doc_date, d.created_at))
    offsets = {o["id"]: o for d in docs for o in svc.list_for(d.id)}
    return {"documents": [ser.document(d) for d in docs], "offsets": list(offsets.values())}


@router.get("/documents/{doc_id}/offsets")
def document_offsets(doc_id: uuid.UUID, user: UserPrincipal = Depends(need("doc.view")), s: Session = SESSION):
    _doc_or_404(s, user, doc_id)
    return {"items": OffsetService(s).list_for(doc_id)}
