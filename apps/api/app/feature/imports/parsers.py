"""ImportParser（開發文件 4.4）：把 Excel 解析成「列」，不碰資料庫。

每種格式一個 parser，新增格式不改呼叫端。欄位名稱容許常見別名，找不到必要欄位就回報錯誤。

【待確認】實際的應付表與薪資表格式。目前以《標準範本》實作（見 README「匯入範本」），
拿到真實檔案後，只需調整 COLUMN_ALIASES 或新增 parser。
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

ID_NUMBER_RE = re.compile(r"^[A-Z][12]\d{8}$")
TAX_ID_RE = re.compile(r"^\d{8}$")
INVOICE_RE = re.compile(r"[A-Z]{2}\d{8}")


@dataclass
class ParsedRow:
    row: int
    values: dict[str, Any]
    warnings: list[dict] = field(default_factory=list)


@dataclass
class ParseResult:
    rows: list[ParsedRow]
    errors: list[dict]
    warnings: list[dict]
    meta: dict[str, Any] = field(default_factory=dict)


def _norm(v: Any) -> str:
    return re.sub(r"\s+", "", str(v)) if v is not None else ""


def _find_header(ws, aliases: dict[str, list[str]], required: list[str], max_scan: int = 15):
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        cells = [_norm(c) for c in row]
        mapping: dict[str, int] = {}
        for key, names in aliases.items():
            for i, c in enumerate(cells):
                # 兩個字的別名只接受完全相同（避免「發票」誤配到「發票金額」）
                if c and any(c == n or (len(n) >= 3 and c.startswith(n)) for n in names):
                    mapping.setdefault(key, i)
                    break
        if all(k in mapping for k in required):
            return r_idx, mapping
    return None, {}


def _to_int(v: Any) -> int | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v).replace(",", "").replace("$", "").strip()
    try:
        return int(round(float(s)))
    except ValueError:
        return None


def _to_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    s = str(v).strip().replace("/", "-").replace(".", "-")
    m = re.match(r"^(\d{2,4})-(\d{1,2})-(\d{1,2})$", s)
    if not m:
        return None
    y = int(m.group(1))
    y = y + 1911 if y < 1000 else y  # 民國年
    return date(y, int(m.group(2)), int(m.group(3)))


def mask_id_number(v: str) -> str:
    return v[:2] + "*****" + v[-3:]


# ---------------------------------------------------------------------- 應付表
AP_ALIASES = {
    "store": ["門市", "店別", "店名"],
    "vendor": ["廠商", "廠商名稱", "供應商", "對象"],
    "vendor_code": ["代號", "統編", "統一編號", "廠商代號"],
    "invoice_no": ["發票號碼", "發票"],
    "date": ["傳票日期", "日期", "發票日期"],
    "memo": ["摘要", "品名", "說明"],
    "amount": ["應付金額", "金額", "含稅金額"],
    "kind": ["類型"],
    "category": ["類別", "類別代碼"],
}


def parse_ap_sheet(data: bytes) -> ParseResult:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    header_row, col = _find_header(ws, AP_ALIASES, ["vendor", "amount"])
    if header_row is None:
        return ParseResult([], [{"row": 0, "code": "NO_HEADER", "message": "找不到「廠商」與「金額」欄"}], [])

    rows: list[ParsedRow] = []
    errors: list[dict] = []
    warnings: list[dict] = []
    section = "cost"
    seen_separator = False
    for r_idx, raw in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        get = lambda k: raw[col[k]] if k in col and col[k] < len(raw) else None  # noqa: E731
        vendor = _norm(get("vendor"))
        joined = "".join(_norm(c) for c in raw if c is not None)
        if not joined:
            continue
        if "成本合計" in joined:
            section, seen_separator = "expense", True
            continue
        if any(k in vendor for k in ("合計", "總計", "小計")) or (not vendor and "合計" in joined):
            continue
        amount = _to_int(get("amount"))
        if not vendor or amount is None:
            errors.append({"row": r_idx, "code": "MISSING_FIELD", "message": "缺廠商或金額"})
            continue
        pr = ParsedRow(row=r_idx, values={
            "store": _norm(get("store")) or None, "vendor": vendor, "amount": amount,
            "memo": str(get("memo")).strip() if get("memo") else None,
            "date": _to_date(get("date")).isoformat() if _to_date(get("date")) else None,
            "kind": {"成本": "cost", "費用": "expense"}.get(_norm(get("kind")), section if "kind" not in col else None),
            "category": _norm(get("category")) or None,
        })
        code = _norm(get("vendor_code"))
        if ID_NUMBER_RE.match(code):
            pr.values["id_masked"] = mask_id_number(code)
            pr.warnings.append({"code": "ID_NUMBER", "message": "代號為身分證字號，已遮蔽"})
        elif TAX_ID_RE.match(code):
            pr.values["tax_id"] = code
        inv_raw = _norm(get("invoice_no"))
        if not inv_raw:
            pr.values["invoice_nos"], pr.values["has_invoice"] = [], False
            pr.warnings.append({"code": "NO_INVOICE", "message": "無發票號碼：進項不能扣抵"})
        elif "沒發票" in inv_raw or "無發票" in inv_raw:
            pr.values["invoice_nos"], pr.values["has_invoice"] = [], False
            pr.warnings.append({"code": "NOTED_NO_INVOICE", "message": "註明沒發票"})
        else:
            nos = INVOICE_RE.findall(inv_raw.upper())
            pr.values["invoice_nos"], pr.values["has_invoice"] = nos or [inv_raw], True
            if "+" in inv_raw or len(nos) > 1:
                pr.warnings.append({"code": "MULTI_INVOICE", "message": f"一格多張發票（{len(nos) or '?'} 張）"})
        rows.append(pr)
        warnings.extend({"row": r_idx, **w} for w in pr.warnings)
    if not seen_separator and "kind" not in col:
        warnings.append({"row": 0, "code": "NO_SEPARATOR", "message": "找不到「成本合計」分隔列，類型改用廠商預設"})
        for r in rows:
            r.values["kind"] = None
    return ParseResult(rows, errors, warnings, {"header_row": header_row})


# ---------------------------------------------------------------------- 發票系統「廠商應付彙總表」
# 一列＝入帳公司 × 門市 × 廠商 × 發票月份 × 入帳月份 的含稅合計；一個檔案涵蓋所有公司。
INVOICE_SUMMARY_ALIASES = {
    "company_name": ["入帳公司名稱", "入帳公司"],
    "company_tax_id": ["公司統編"],
    "store": ["門市名稱"],
    "store_code": ["門市代碼"],
    "invoice_month": ["發票月份"],
    "period": ["入帳月份"],
    "vendor": ["廠商名稱"],
    "vendor_tax_id": ["統一編號", "廠商統編"],
    "invoice_count": ["發票張數"],
    "amount": ["含稅應付總計", "應付總計", "含稅金額"],
    "due": ["預計付款日"],
}
_MONTH_RE = re.compile(r"^(\d{4})-(\d{1,2})$")


def _month(v: Any) -> str | None:
    if isinstance(v, (datetime, date)):
        return f"{v.year:04d}-{v.month:02d}"
    m = _MONTH_RE.match(_norm(v).replace("/", "-"))
    if not m or not 1 <= int(m.group(2)) <= 12:
        return None
    return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"


def looks_like_invoice_summary(data: bytes) -> bool:
    try:
        ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).worksheets[0]
    except Exception:  # noqa: BLE001
        return False
    row, _ = _find_header(ws, INVOICE_SUMMARY_ALIASES, ["company_tax_id", "vendor", "amount", "period"])
    return row is not None


def parse_invoice_summary(data: bytes) -> ParseResult:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    header_row, col = _find_header(ws, INVOICE_SUMMARY_ALIASES, ["company_tax_id", "vendor", "amount", "period"])
    if header_row is None:
        return ParseResult([], [{"row": 0, "code": "NO_HEADER",
                                 "message": "找不到「公司統編」「廠商名稱」「入帳月份」「含稅應付總計」欄"}], [])
    rows: list[ParsedRow] = []
    errors: list[dict] = []
    warnings: list[dict] = []
    for r_idx, raw in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        get = lambda k: raw[col[k]] if k in col and col[k] < len(raw) else None  # noqa: E731
        if not any(v is not None and str(v).strip() for v in raw):
            continue
        vendor = str(get("vendor") or "").strip()
        if any(k in vendor for k in ("合計", "總計", "小計")):
            continue
        amount = _to_int(get("amount"))
        period = _month(get("period"))
        company_tax_id = _norm(get("company_tax_id"))
        missing = [n for n, ok in (("廠商名稱", vendor), ("含稅應付總計", amount is not None),
                                   ("入帳月份", period), ("公司統編", TAX_ID_RE.match(company_tax_id))) if not ok]
        if missing:
            errors.append({"row": r_idx, "code": "MISSING_FIELD", "message": f"缺少或格式錯誤：{'、'.join(missing)}"})
            continue
        if amount == 0:
            continue
        pr = ParsedRow(row=r_idx, values={
            "company_tax_id": company_tax_id, "company_name": _norm(get("company_name")) or None,
            "store": str(get("store") or "").strip() or None, "store_code": _norm(get("store_code")) or None,
            "period": period, "vendor": vendor, "amount": amount,
            "invoice_count": _to_int(get("invoice_count")) or 0,
        })
        inv_month = _month(get("invoice_month"))
        if inv_month is None:
            pr.values["invoice_month"] = period
            pr.warnings.append({"code": "BAD_INVOICE_MONTH",
                                "message": f"發票月份「{get('invoice_month')}」格式錯誤，以入帳月份 {period} 代替"})
        else:
            pr.values["invoice_month"] = inv_month
        due_m = re.search(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", str(get("due") or ""))
        pr.values["due"] = f"{due_m.group(1)}-{int(due_m.group(2)):02d}-{int(due_m.group(3)):02d}" if due_m else None
        code = _norm(get("vendor_tax_id")).upper()
        if ID_NUMBER_RE.match(code):
            pr.values["id_masked"] = mask_id_number(code)
            pr.warnings.append({"code": "ID_NUMBER", "message": "廠商統編欄是身分證字號（個人），已遮蔽；視為無發票"})
        elif TAX_ID_RE.match(code):
            pr.values["tax_id"] = code
        else:
            pr.warnings.append({"code": "NO_TAX_ID", "message": "廠商沒有統一編號，只能用名稱比對"})
        rows.append(pr)
        warnings.extend({"row": r_idx, **w} for w in pr.warnings)
    return ParseResult(rows, errors, warnings, {"header_row": header_row})


# ---------------------------------------------------------------------- 薪資表（門市彙總）
PAYROLL_ALIASES = {
    "store": ["門市", "店別", "店名", "部門"],
    "P01": ["薪資", "本薪", "薪資合計"],
    "P02": ["加班與獎金", "加班費", "獎金"],
    "P03_labor": ["勞保雇主", "勞保公司負擔", "勞保"],
    "P03_health": ["健保雇主", "健保公司負擔", "健保"],
    "P04": ["勞退", "勞退提繳", "勞退6%"],
}


def parse_payroll(data: bytes) -> ParseResult:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    header_row, col = _find_header(ws, PAYROLL_ALIASES, ["store", "P01"])
    if header_row is None:
        return ParseResult([], [{"row": 0, "code": "NO_HEADER", "message": "找不到「門市」與「薪資」欄"}], [])
    rows, errors = [], []
    for r_idx, raw in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        store = _norm(raw[col["store"]]) if col["store"] < len(raw) else ""
        if not store or any(k in store for k in ("合計", "總計")):
            continue
        vals: dict[str, int] = {}
        for key in ("P01", "P02", "P03_labor", "P03_health", "P04"):
            if key in col:
                vals[key] = _to_int(raw[col[key]]) or 0
        if not any(vals.values()):
            errors.append({"row": r_idx, "code": "EMPTY_ROW", "message": f"{store} 沒有金額"})
            continue
        rows.append(ParsedRow(row=r_idx, values={"store": store, **vals}))
    return ParseResult(rows, errors, [])


# ---------------------------------------------------------------- POS 每日營業報表
# POS 下載的「Sales Summary」：每天一列；「內用」「外帶」分頁各有每種收款方式的金額。
POS_SHEETS = ("Sales Summary", "內用", "外帶")
POS_OTHER_COLS = ("預訂單訂金", "預定單訂金退回", "禮券退款", "進位差異", "小費", "服務費", "送遞附加費")


def looks_like_pos_daily(data: bytes) -> bool:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return all(n in wb.sheetnames for n in POS_SHEETS)
    except Exception:  # noqa: BLE001
        return False


def _pos_sheet(ws) -> tuple[dict[str, int], dict[str, int], list[tuple]]:
    """回傳（欄名→欄位, 收款方式→金額欄位, 資料列）。收款方式在「日期」上一列，金額是名稱下方那欄右邊的 NT$。"""
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and _norm(r[0]) == "日期")
    header = [_norm(c) for c in rows[hi]]
    above = rows[hi - 1] if hi > 0 else ()
    cols = {h: i for i, h in enumerate(header) if h}
    tenders: dict[str, int] = {}
    for i, name in enumerate(above):
        if name and i + 1 < len(header) and header[i + 1].upper() == "NT$":
            tenders[str(name).strip()] = i + 1
    data = [r for r in rows[hi + 1:] if r and r[0] and _to_date(str(r[0])[:10])]
    return cols, tenders, data


def parse_pos_daily(data: bytes) -> ParseResult:
    """一列＝一天＋一種收款方式（現金、信用卡、LINE Pay…），金額拆成內用、外帶。"""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    missing = [n for n in POS_SHEETS if n not in wb.sheetnames]
    if missing:
        return ParseResult([], [{"row": 0, "code": "MISSING_SHEET", "message": f"缺少分頁：{'、'.join(missing)}"}], [])
    summary, dine, take = (_pos_sheet(wb[n]) for n in POS_SHEETS)
    by_day: dict[date, dict[str, Any]] = {}
    for label, (cols, tenders, rows) in (("dine", dine), ("take", take)):
        for r in rows:
            d = _to_date(str(r[0])[:10])
            day = by_day.setdefault(d, {"dine": {}, "take": {}})
            for t, ci in tenders.items():
                day[label][t] = _to_int(r[ci]) or 0
    out, warnings, errors = [], [], []
    s_cols, s_tenders, s_rows = summary
    for idx, r in enumerate(s_rows, start=1):
        d = _to_date(str(r[0])[:10])
        day = by_day.get(d, {"dine": {}, "take": {}})
        others = {c: _to_int(r[s_cols[c]]) or 0 for c in POS_OTHER_COLS if c in s_cols}
        nonzero = {k: v for k, v in others.items() if v}
        row_w = [{"code": "POS_OTHER", "message": f"{d} 有 {'、'.join(f'{k} {v:,}' for k, v in nonzero.items())}，"
                  "目前沒有入帳，請確認"}] if nonzero else []
        for t, ci in s_tenders.items():
            total = _to_int(r[ci]) or 0
            if not total:
                continue
            dn, tk = day["dine"].get(t, 0), day["take"].get(t, 0)
            w = list(row_w)
            if dn + tk != total:
                w.append({"code": "POS_SPLIT_MISMATCH", "message": f"{d} {t}：內用 {dn:,}＋外帶 {tk:,} ≠ 總表 {total:,}，差額記外帶"})
                tk = total - dn
            out.append(ParsedRow(row=idx + 4, values={"date": d, "tender": t, "total": total, "dine_in": dn, "takeout": tk},
                                 warnings=w))
            row_w = []  # 同一天的其他欄位警告只提一次
    if not out:
        errors.append({"row": 0, "code": "NO_ROWS", "message": "找不到每日營業資料"})
    return ParseResult(out, errors, warnings)


PARSERS = {"ap_sheet": parse_ap_sheet, "payroll": parse_payroll, "invoice_summary": parse_invoice_summary,
           "pos_daily": parse_pos_daily}
