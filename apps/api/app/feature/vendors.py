"""廠商主檔（開發文件 v0.6 5.16、5.17）：統編檢查、搜尋、單筆新增、Excel 批次匯入。

- 統編是預付沖抵、折讓對單的比對依據，所以新增廠商一定要有統編，或明確勾「沒有統編」（個人、小販）。
- 批次匯入分兩段：先預覽（不寫入），再確認；確認時可指定「名稱不同」的列要不要改名。
- 匯入只新增或補資料，不刪除任何廠商。
- 類別對應（vendor_map.py）仍是另一張表；這裡也接受「預設類別」欄，有填才設定。
"""
from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.errors import NotFound, TaxIdExists, TaxIdRequired, ValidationFailed
from app.feature.imports.parsers import ID_NUMBER_RE, _find_header, _norm, mask_id_number
from app.feature.models import Counterparty, Document, DocumentOffset
from app.feature.repositories import CatalogRepository
from app.seams.company import get_directory

TAX_WEIGHTS = (1, 2, 1, 2, 1, 2, 4, 1)
MASKED_SYS_RE = re.compile(r"^[A-Z][12]\*{5}\d{3}$")      # 系統格式：T1*****011
MASKED_ALT_RE = re.compile(r"^([A-Z][12])\d\*{4}(\d{3})$")  # 整理表格式：T12****011
COMPANY_SUFFIXES = ("股份有限公司", "有限公司", "（股）公司", "(股)公司", "（股）", "(股)", "公司")
STOPPED = {"停用", "歇業", "非營業中", "合併解散", "廢止", "查無資料"}

ALIASES = {
    "name": ["廠商名稱", "廠商", "全稱", "名稱"],
    "tax_id": ["廠商統編", "統一編號", "統編"],
    "no_tax_id": ["沒有統編"],
    "masked": ["遮蔽身分證"],
    "short": ["簡稱", "別名"],
    "category": ["預設類別", "類別代碼"],
    "keywords": ["付款關鍵字", "關鍵字"],
    "status": ["狀態"],
    "intercompany": ["集團內"],
}


def valid_tax_id(t: str | None) -> bool:
    """營利事業統一編號檢查碼（財政部規則；2023 起可被 5 整除，舊號被 10 整除也符合）。"""
    if not t or not re.fullmatch(r"\d{8}", t):
        return False
    s = 0
    for d, w in zip(t, TAX_WEIGHTS):
        p = int(d) * w
        s += p // 10 + p % 10
    return s % 5 == 0 or (t[6] == "7" and (s + 1) % 5 == 0)


def normalize_name(n: str) -> str:
    n = re.sub(r"\s+", "", n or "").replace("故份", "股份")
    for suf in COMPANY_SUFFIXES:
        n = n.replace(suf, "")
    return re.sub(r"[()（）\-－・]", "", n)


def normalize_masked(v: str) -> str | None:
    v = (v or "").strip().upper()
    if MASKED_SYS_RE.match(v):
        return v
    m = MASKED_ALT_RE.match(v)
    if m:
        return m.group(1) + "*****" + m.group(2)
    if ID_NUMBER_RE.match(v):
        return mask_id_number(v)
    return None


@dataclass
class PreviewRow:
    row: int
    action: str  # create / fill_tax_id / exists / name_differs / error / skip
    name: str
    tax_id: str | None
    message: str
    counterparty_id: str | None = None
    current_name: str | None = None
    inactive: bool = False


@dataclass
class ImportReport:
    rows: list[PreviewRow] = field(default_factory=list)
    counts: dict[str, int] = field(default_factory=dict)
    committed: bool = False

    def to_dict(self) -> dict[str, Any]:
        return {"committed": self.committed, "counts": self.counts,
                "rows": [r.__dict__ for r in self.rows]}


class VendorMasterService:
    def __init__(self, session: Session):
        self.s = session
        self.catalog = CatalogRepository(session)

    # ------------------------------------------------------------------ 搜尋（5.16）
    def search(self, q: str, limit: int = 10, roles: str | None = "vendor",
               company_id: str | None = None) -> list[dict[str, Any]]:
        q = (q or "").strip()
        if not q:
            return []
        rows = list(self.s.scalars(select(Counterparty).where(Counterparty.is_active.is_(True))))
        if roles:
            rows = [c for c in rows if roles in (c.roles or ["vendor"]) or c.is_intercompany]
        hits: list[tuple[int, Counterparty, str]] = []
        digits = q.isdigit()
        nq = normalize_name(q)
        for c in rows:
            if digits:
                if c.tax_id == q:
                    hits.append((0, c, "tax_id"))
                elif c.tax_id and c.tax_id.startswith(q):
                    hits.append((1, c, "tax_id"))
                continue
            names = [c.name, *(c.aliases or [])]
            if any(q == n for n in names):
                hits.append((0, c, "name"))
            elif any(q in n for n in names) or (nq and any(nq in normalize_name(n) for n in names)):
                hits.append((2, c, "name"))
            elif any(k and (k in q or q in k) for k in (c.keywords or [])):
                hits.append((3, c, "keyword"))
        hits.sort(key=lambda h: (h[0], len(h[1].name), h[1].name))
        hits = hits[:limit]
        open_pre = self.open_prepayment_amounts({h[1].id for h in hits}, company_id)
        similar = len({normalize_name(h[1].name)[:2] for h in hits}) < len(hits)
        return [{"id": str(c.id), "name": c.name, "tax_id": c.tax_id, "id_masked": c.id_masked,
                 "no_tax_id": c.no_tax_id, "aliases": c.aliases or [], "is_intercompany": c.is_intercompany,
                 "intercompany_company_id": c.intercompany_company_id,
                 "default_category_code": c.default_category_code, "match": m,
                 "open_prepayment_amount": open_pre.get(c.id, 0), "similar_names": similar and len(hits) > 1}
                for _, c, m in hits]

    def open_prepayment_amounts(self, cp_ids: set[uuid.UUID], company_id: str | None) -> dict[uuid.UUID, int]:
        if not cp_ids:
            return {}
        used = (select(DocumentOffset.from_document_id, func.sum(DocumentOffset.amount).label("u"))
                .where(DocumentOffset.reversed_at.is_(None)).group_by(DocumentOffset.from_document_id).subquery())
        stmt = (select(Document.counterparty_id, func.sum(Document.settle_amount - func.coalesce(used.c.u, 0)))
                .outerjoin(used, used.c.from_document_id == Document.id)
                .where(Document.doc_type == "prepayment", Document.status == "posted",
                       Document.counterparty_id.in_(cp_ids)).group_by(Document.counterparty_id))
        if company_id:
            d = get_directory()
            ent = d.legal_entity(company_id)
            cids = [c.id for c in d.list_companies() if d.legal_entity(c.id) == ent]
            stmt = stmt.where(Document.company_id.in_(cids))
        return {cid: int(v or 0) for cid, v in self.s.execute(stmt) if v}

    # ------------------------------------------------------------------ 清單（S13）
    def list(self, q: str | None = None, missing_tax_id: bool = False, include_inactive: bool = False) -> list[dict]:
        stmt = select(Counterparty).order_by(Counterparty.name)
        if not include_inactive:
            stmt = stmt.where(Counterparty.is_active.is_(True))
        out = []
        for c in self.s.scalars(stmt):
            if "vendor" not in (c.roles or ["vendor"]) and not c.is_intercompany:
                continue
            lacks = not c.tax_id and not c.no_tax_id and not c.id_masked
            if missing_tax_id and not lacks:
                continue
            if q and q not in c.name and q not in (c.tax_id or "") and not any(q in a for a in c.aliases or []):
                continue
            out.append({"id": str(c.id), "name": c.name, "tax_id": c.tax_id, "id_masked": c.id_masked,
                        "no_tax_id": c.no_tax_id, "missing_tax_id": lacks, "aliases": c.aliases or [],
                        "keywords": c.keywords or [], "default_category_code": c.default_category_code,
                        "payment_rule": c.payment_rule, "is_intercompany": c.is_intercompany,
                        "is_active": c.is_active})
        return out

    # ------------------------------------------------------------------ 單筆新增／修改
    def create(self, *, name: str, tax_id: str | None, no_tax_id: bool = False, id_masked: str | None = None,
               aliases: list[str] | None = None, keywords: list[str] | None = None,
               default_category_code: str | None = None, payment_rule: str = "next_month_30",
               is_intercompany: bool = False, intercompany_company_id: str | None = None) -> Counterparty:
        name = (name or "").strip()
        if not name:
            raise ValidationFailed("廠商名稱必填", [{"field": "name", "message": "必填"}])
        tax_id = (tax_id or "").strip() or None
        if tax_id:
            if not valid_tax_id(tax_id):
                raise ValidationFailed("統編格式或檢查碼不正確", [{"field": "tax_id", "message": "檢查碼不符"}])
            hit = self.catalog.counterparty_by_tax_id(tax_id)
            if hit:
                raise TaxIdExists(f"這個統編已經是「{hit.name}」", [{"field": "tax_id", "message": str(hit.id)}])
        elif not no_tax_id:
            raise TaxIdRequired("請填廠商統編；沒有統編請勾選「沒有統編」",
                                [{"field": "tax_id", "message": "必填"}])
        if default_category_code and self.catalog.category(default_category_code) is None:
            raise ValidationFailed("預設類別不存在")
        group = get_directory().get_by_tax_id(tax_id) if tax_id else None
        return self.catalog.add_counterparty(
            name=name, tax_id=tax_id, no_tax_id=bool(no_tax_id and not tax_id),
            id_masked=normalize_masked(id_masked) if id_masked else None, roles=["vendor"],
            aliases=[a for a in (aliases or []) if a and a != name], keywords=keywords or [],
            default_category_code=default_category_code or None, payment_rule=payment_rule or "next_month_30",
            is_intercompany=is_intercompany or group is not None,
            intercompany_company_id=intercompany_company_id or (group.id if group else None))

    def update(self, cp_id: uuid.UUID, **changes) -> Counterparty:
        cp = self.catalog.counterparty(cp_id)
        if cp is None:
            raise NotFound("廠商不存在")
        if "tax_id" in changes and changes["tax_id"] != cp.tax_id:
            t = (changes["tax_id"] or "").strip() or None
            if t:
                if not valid_tax_id(t):
                    raise ValidationFailed("統編格式或檢查碼不正確", [{"field": "tax_id", "message": "檢查碼不符"}])
                hit = self.catalog.counterparty_by_tax_id(t)
                if hit and hit.id != cp.id:
                    raise TaxIdExists(f"這個統編已經是「{hit.name}」", [{"field": "tax_id", "message": str(hit.id)}])
                cp.no_tax_id = False
            cp.tax_id = t
        if changes.get("name") and changes["name"].strip() != cp.name:
            old = cp.name
            cp.name = changes["name"].strip()
            cp.aliases = [a for a in {*(cp.aliases or []), old} if a != cp.name]
        for k in ("aliases", "keywords", "default_category_code", "payment_rule", "is_active", "no_tax_id"):
            if k in changes and changes[k] is not None:
                setattr(cp, k, changes[k])
        if cp.default_category_code and self.catalog.category(cp.default_category_code) is None:
            raise ValidationFailed("預設類別不存在")
        self.s.flush()
        return cp

    # ------------------------------------------------------------------ Excel 批次匯入（5.17）
    def import_xlsx(self, data: bytes, *, commit: bool, decisions: dict[int, str] | None = None) -> ImportReport:
        """decisions：{列號: "update_name" | "keep_name" | "skip"}，只用在「名稱不同」的列。"""
        decisions = decisions or {}
        ws, header_row, col = self._open(data)
        idx = self._index()
        rep = ImportReport()
        seen_tax: set[str] = set()
        for r_idx, raw in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            get = lambda k: raw[col[k]] if k in col and col[k] < len(raw) else None  # noqa: E731
            name = str(get("name") or "").strip()
            tax = _norm(get("tax_id")).upper()
            if tax.isdigit() and len(tax) < 8:
                tax = tax.zfill(8)  # Excel 把統編當數字，開頭的 0 不見了
            if not name and not tax:
                continue
            no_tax = _norm(get("no_tax_id")).upper() in ("Y", "是", "V", "TRUE", "1") or bool(ID_NUMBER_RE.match(tax))
            masked = normalize_masked(_norm(get("masked")) or (tax if ID_NUMBER_RE.match(tax) else ""))
            if ID_NUMBER_RE.match(tax):
                tax = ""  # 身分證不存，只留遮蔽號碼
            short = [a for a in re.split(r"[、,，/]", str(get("short") or "")) if a.strip()]
            kws = [k for k in re.split(r"[、,，/]", str(get("keywords") or "")) if k.strip()]
            category = _norm(get("category")).upper()[:3] or None
            inactive = _norm(get("status")) in STOPPED
            if not name:
                rep.rows.append(PreviewRow(r_idx, "error", name, tax or None, "廠商名稱空白"))
                continue
            if tax and not valid_tax_id(tax):
                rep.rows.append(PreviewRow(r_idx, "error", name, tax, "統編格式或檢查碼不正確，這一列不匯入"))
                continue
            if category and self.catalog.category(category) is None:
                rep.rows.append(PreviewRow(r_idx, "error", name, tax or None, f"預設類別 {category} 不存在"))
                continue
            if not tax and not no_tax:
                rep.rows.append(PreviewRow(r_idx, "error", name, None, "沒有統編；若是個人或小販，請在「沒有統編」填 Y"))
                continue
            if tax and tax in seen_tax:
                rep.rows.append(PreviewRow(r_idx, "skip", name, tax, "檔案中統編重複，以第一次出現的為準"))
                continue
            if tax:
                seen_tax.add(tax)

            cp = idx["tax"].get(tax) if tax else None
            if cp is not None:
                if cp.name == name or name in (cp.aliases or []):
                    row = PreviewRow(r_idx, "exists", name, tax, "已存在", str(cp.id), cp.name, inactive)
                    rep.rows.append(row)
                    if commit:
                        self._merge_extras(cp, short, kws, category, inactive)
                    continue
                choice = decisions.get(r_idx, "keep_name")
                rep.rows.append(PreviewRow(r_idx, "name_differs", name, tax,
                                           f"主檔是「{cp.name}」", str(cp.id), cp.name, inactive))
                if commit and choice != "skip":
                    if choice == "update_name":
                        cp.aliases = [a for a in {*(cp.aliases or []), cp.name} if a != name]
                        cp.name = name
                    elif name not in (cp.aliases or []):
                        cp.aliases = [*(cp.aliases or []), name]
                    self._merge_extras(cp, short, kws, category, inactive)
                continue

            # 統編沒有：名稱（含別名、付款關鍵字）對到還沒有統編的廠商 → 補統編
            cand = self._match_untaxed(idx, name, short, masked)
            if cand is not None and not tax and cand.no_tax_id:
                rep.rows.append(PreviewRow(r_idx, "exists", name, None, "已存在（沒有統編）", str(cand.id), cand.name,
                                           inactive))
                if commit:
                    self._merge_extras(cand, short, kws, category, inactive)
                continue
            if cand is not None:
                msg = "補上統編" if tax else "標記為沒有統編"
                rep.rows.append(PreviewRow(r_idx, "fill_tax_id", name, tax or None, msg, str(cand.id), cand.name,
                                           inactive))
                if commit:
                    if tax:
                        cand.tax_id = tax
                        cand.no_tax_id = False
                        idx["tax"][tax] = cand
                    else:
                        cand.no_tax_id = True
                        cand.id_masked = cand.id_masked or masked
                    if cand.name != name:
                        cand.aliases = [a for a in {*(cand.aliases or []), cand.name} if a != name]
                        cand.name = name
                    self._merge_extras(cand, short, kws, category, inactive)
                continue
            if not tax and idx["name"].get(name) is not None:
                ex = idx["name"][name]
                rep.rows.append(PreviewRow(r_idx, "exists", name, None, "已存在", str(ex.id), ex.name, inactive))
                continue

            rep.rows.append(PreviewRow(r_idx, "create", name, tax or None,
                                       "新增" + ("（停用）" if inactive else ""), inactive=inactive))
            if commit:
                group = get_directory().get_by_tax_id(tax) if tax else None
                cp = Counterparty(id=uuid.uuid4(), name=name, tax_id=tax or None, no_tax_id=not tax,
                                  id_masked=masked if not tax else None, roles=["vendor"],
                                  aliases=[a for a in short if a != name], keywords=kws,
                                  default_category_code=category, is_intercompany=group is not None,
                                  intercompany_company_id=group.id if group else None,
                                  payment_rule="next_month_30", is_active=not inactive, issues_invoice=bool(tax),
                                  bank_accounts=[])
                self.s.add(cp)
                if tax:
                    idx["tax"][tax] = cp
                idx["name"].setdefault(name, cp)
        if commit:
            self.s.flush()
        rep.committed = commit
        for r in rep.rows:
            rep.counts[r.action] = rep.counts.get(r.action, 0) + 1
        return rep

    # ------------------------------------------------------------------ 內部
    def _open(self, data: bytes):
        try:
            wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        except Exception as e:  # noqa: BLE001
            raise ValidationFailed(f"讀不到這個檔案（{type(e).__name__}），請上傳 .xlsx") from e
        sheets = sorted(wb.worksheets, key=lambda w: 0 if w.title == "匯入名單" else 1)
        for ws in sheets:
            h, c = _find_header(ws, ALIASES, ["name", "tax_id"])
            if h is not None:
                return ws, h, c
        raise ValidationFailed("找不到「廠商名稱」「廠商統編」兩欄；請用匯入範本或整理好的「匯入名單」工作表")

    def _index(self) -> dict:
        idx: dict = {"tax": {}, "name": {}, "untaxed": []}
        for cp in self.s.scalars(select(Counterparty)):
            if cp.tax_id:
                idx["tax"][cp.tax_id] = cp
            else:
                idx["untaxed"].append(cp)
            for n in [cp.name, *(cp.aliases or [])]:
                idx["name"].setdefault(n, cp)
        return idx

    @staticmethod
    def _match_untaxed(idx: dict, name: str, short: list[str], masked: str | None) -> Counterparty | None:
        names = {name, *short}
        nn = normalize_name(name)
        exact = [c for c in idx["untaxed"] if names & {c.name, *(c.aliases or [])}]
        if len(exact) == 1:
            return exact[0]
        if masked:
            m = [c for c in idx["untaxed"] if c.id_masked == masked]
            if len(m) == 1:
                return m[0]
        fuzzy = [c for c in idx["untaxed"] if nn and normalize_name(c.name) == nn]
        if len(fuzzy) == 1:
            return fuzzy[0]
        kw = [c for c in idx["untaxed"] if any(k and len(k) >= 2 and k in name for k in (c.keywords or []))]
        return kw[0] if len(kw) == 1 else None

    @staticmethod
    def _merge_extras(cp: Counterparty, short: list[str], kws: list[str], category: str | None,
                      inactive: bool) -> None:
        cp.aliases = [*(cp.aliases or []), *[a for a in short if a != cp.name and a not in (cp.aliases or [])]]
        cp.keywords = [*(cp.keywords or []), *[k for k in kws if k not in (cp.keywords or [])]]
        if category and not cp.default_category_code:
            cp.default_category_code = category  # 只補空白，不覆蓋已設定的類別
        if inactive:
            cp.is_active = False
