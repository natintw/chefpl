"""廠商類別對應表：廠商（以統一編號為準）→ 管理類別（mgmt_categories.code）。

發票系統只給廠商與金額，沒有科目，所以靠這張表決定每張發票進損益表的哪一行。
實際存在 counterparties.default_category_code，分類引擎的「廠商預設」層會用它（見 classification.py）。

- 以統一編號比對；個人廠商用遮蔽後的身分證字號＋姓名；都沒有才用名稱（含別名）。
- 設定或修改類別時，這個廠商「等待分類」的草稿會一起入帳，待確認同步關掉。
- 匯入檔可以順便修正廠商名稱：統編相同、名稱不同時，以檔案的名稱為準，舊名稱留作別名。
"""
from __future__ import annotations

import io
import re
import uuid
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import case, func, select
from sqlalchemy.orm import Session

from app.errors import NotFound, ValidationFailed
from app.feature.imports.parsers import ID_NUMBER_RE, TAX_ID_RE, _find_header, _norm, mask_id_number
from app.feature.models import Counterparty, Document, ReviewItem
from app.feature.repositories import CatalogRepository, OpsRepository

MASKED_RE = re.compile(r"^[A-Z][12]\*{5}\d{3}$")
# 標準格式（一個工作表、四欄）：統一編號｜廠商名稱｜類別代碼｜類別名稱
HEADERS = ["統一編號", "廠商名稱", "類別代碼", "類別名稱"]
ALIASES = {
    "tax_id": ["統一編號", "廠商統編", "統編"],
    "name": ["廠商名稱", "廠商"],
    "category": ["類別代碼"],
    "category_name": ["類別名稱"],
}
VENDOR_NATURES = {"expense", "balance"}


@dataclass
class ImportResult:
    updated: int = 0
    created: int = 0
    unchanged: int = 0
    skipped_blank: int = 0
    applied_drafts: int = 0
    errors: list[dict[str, Any]] = field(default_factory=list)


class VendorMapService:
    def __init__(self, session: Session):
        self.s = session
        self.catalog = CatalogRepository(session)
        self.ops = OpsRepository(session)

    # ------------------------------------------------------------------ 查詢
    def list(self, q: str | None = None, only_unmapped: bool = False) -> list[dict[str, Any]]:
        stats = {row.counterparty_id: row for row in self.s.execute(
            select(Document.counterparty_id, func.count().label("n"),
                   func.coalesce(func.sum(Document.settle_amount), 0).label("amount"),
                   func.sum(case((Document.status == "draft", 1), else_=0)).label("drafts"))
            .where(Document.status != "void", Document.doc_type == "payable")
            .group_by(Document.counterparty_id))}
        out = []
        for cp in self.s.scalars(select(Counterparty).where(Counterparty.is_active.is_(True))):
            if "vendor" not in (cp.roles or ["vendor"]):
                continue
            if only_unmapped and cp.default_category_code:
                continue
            if q and q not in cp.name and q not in (cp.tax_id or "") and not any(q in a for a in cp.aliases or []):
                continue
            st = stats.get(cp.id)
            cat = self.catalog.category(cp.default_category_code) if cp.default_category_code else None
            out.append({"id": str(cp.id), "name": cp.name, "tax_id": cp.tax_id, "id_masked": cp.id_masked,
                        "aliases": cp.aliases or [], "category_code": cp.default_category_code,
                        "category_name": cat.name if cat else None, "is_intercompany": cp.is_intercompany,
                        "doc_count": int(st.n) if st else 0, "amount": int(st.amount) if st else 0,
                        "drafts": int(st.drafts or 0) if st else 0})
        out.sort(key=lambda r: (r["category_code"] is not None, -r["amount"], r["name"]))
        return out

    # ------------------------------------------------------------------ 單筆設定
    def set(self, cp_id: uuid.UUID, category_code: str | None, actor: str, name: str | None = None) -> dict[str, Any]:
        cp = self.catalog.counterparty(cp_id)
        if cp is None:
            raise NotFound("廠商不存在")
        if category_code:
            self._check_category(category_code)
        if name and name.strip() and name.strip() != cp.name:
            self._rename(cp, name.strip())
        cp.default_category_code = category_code or None
        applied = self._apply_to_drafts(cp, actor) if category_code else 0
        self._close_new_vendor_review(cp, actor)
        return {"id": str(cp.id), "name": cp.name, "category_code": cp.default_category_code, "applied_drafts": applied}

    # ------------------------------------------------------------------ Excel 匯入／匯出
    def import_xlsx(self, data: bytes, actor: str) -> ImportResult:
        try:
            ws = load_workbook(io.BytesIO(data), data_only=True, read_only=True).worksheets[0]
        except Exception as e:  # noqa: BLE001
            raise ValidationFailed(f"讀不到這個檔案（{type(e).__name__}），請上傳 .xlsx") from e
        header_row, col = None, {}
        for need in (["tax_id", "category"], ["tax_id", "category_name"], ["name", "category"], ["name", "category_name"]):
            h, c = _find_header(ws, ALIASES, need)
            if h is not None and (header_row is None or h < header_row):
                header_row, col = h, c
        if header_row is None:
            raise ValidationFailed("格式為四欄：" + "｜".join(HEADERS) + "（類別代碼、類別名稱至少要有一欄）")
        by_name = {c.name: c.code for c in self.catalog.categories.values()}
        res = ImportResult()
        for r_idx, raw in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            get = lambda k: raw[col[k]] if k in col and col[k] < len(raw) else None  # noqa: E731
            name = str(get("name") or "").strip()
            code_raw = _norm(get("tax_id")).upper()
            if code_raw.isdigit() and len(code_raw) < 8:  # Excel 把統編當數字，前面的 0 不見了
                code_raw = code_raw.zfill(8)
            category = _norm(get("category")).upper()
            cat_name = _norm(get("category_name"))
            if cat_name:
                named = by_name.get(cat_name)
                if not category and named:
                    category = named  # 只填類別名稱也可以
                elif category and named and named != category:
                    res.errors.append({"row": r_idx, "message": f"類別代碼 {category} 與類別名稱「{cat_name}」不一致"})
                    continue
                elif not category and not named:
                    res.errors.append({"row": r_idx, "message": f"找不到類別名稱「{cat_name}」"})
                    continue
            if not name and not code_raw:
                continue
            tax_id = code_raw if TAX_ID_RE.match(code_raw) else None
            masked = mask_id_number(code_raw) if ID_NUMBER_RE.match(code_raw) else (
                code_raw if MASKED_RE.match(code_raw) else None)
            cp = self._find(tax_id, masked, name)
            if not category:  # 類別空白：不改類別，但統編對得到時仍可順便更正名稱
                if cp is not None and name and (tax_id or masked) and name != cp.name:
                    self._rename(cp, name)
                res.skipped_blank += 1
                continue
            try:
                self._check_category(category)
            except ValidationFailed as e:
                res.errors.append({"row": r_idx, "message": e.message})
                continue
            if cp is None:
                if not name:
                    res.errors.append({"row": r_idx, "message": f"找不到統編 {code_raw} 的廠商，新增時要填廠商名稱"})
                    continue
                cp = Counterparty(id=uuid.uuid4(), name=name, tax_id=tax_id, id_masked=masked, roles=["vendor"],
                                  default_category_code=category, aliases=[], keywords=[], is_intercompany=False,
                                  is_active=True)
                self.s.add(cp)  # 最後一次寫入
                self._add_to_index(self._index(), cp)
                res.created += 1
                continue
            if name and (tax_id or masked) and name != cp.name:
                self._rename(cp, name)
            if cp.default_category_code == category:
                res.unchanged += 1
            else:
                cp.default_category_code = category
                res.updated += 1
            res.applied_drafts += self._apply_to_drafts(cp, actor)
            self._close_new_vendor_review(cp, actor)
        self.s.flush()
        return res

    def export_xlsx(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "廠商類別"
        ws.append(HEADERS)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E8EEF6")
        for r in self.list():
            if r["is_intercompany"]:
                continue  # 集團內公司固定 C99，不放進對應表
            ws.append([r["tax_id"] or r["id_masked"] or "", r["name"], r["category_code"] or "", r["category_name"] or ""])
        for row in ws.iter_rows(min_row=2, max_col=1):
            row[0].number_format = "@"  # 統編存成文字，保留開頭的 0
        for letter, width in zip("ABCD", (12, 36, 10, 14)):
            ws.column_dimensions[letter].width = width
        ws.freeze_panes = "A2"
        codes = ",".join(c.code for c in self.catalog.categories.values() if c.nature in VENDOR_NATURES and c.is_active)
        if len(codes) <= 255:  # 類別代碼欄的下拉選單（Excel 清單上限 255 字）
            dv = DataValidation(type="list", formula1=f'"{codes}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"C2:C{max(ws.max_row, 2) + 200}")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------ 內部
    def _check_category(self, code: str) -> None:
        cat = self.catalog.category(code)
        if cat is None:
            raise ValidationFailed(f"類別代碼 {code} 不存在")
        if cat.nature not in VENDOR_NATURES:
            raise ValidationFailed(f"{code} {cat.name} 是收入類別，廠商（應付）不能用")

    # ----- 批次用的快取：上傳一百多列時，不要每列都查資料庫 -----
    def _index(self) -> dict:
        if not hasattr(self, "_idx"):
            idx: dict = {"tax": {}, "masked": {}, "name": {}}
            for cp in self.s.scalars(select(Counterparty)):
                self._add_to_index(idx, cp)
            self._idx = idx
        return self._idx

    @staticmethod
    def _add_to_index(idx: dict, cp: Counterparty) -> None:
        if cp.tax_id:
            idx["tax"].setdefault(cp.tax_id, cp)
        if cp.id_masked:
            idx["masked"].setdefault((cp.id_masked, cp.name), cp)
            idx["masked"].setdefault((cp.id_masked, None), cp)
        for n in [cp.name, *(cp.aliases or [])]:
            idx["name"].setdefault(n, cp)

    def _drafts_by_vendor(self) -> dict:
        """所有等待分類的草稿（含明細），一次讀進來：{廠商 ID: {單據 ID: 單據}}。"""
        if not hasattr(self, "_drafts"):
            self._drafts: dict = {}
            for d in self.s.scalars(select(Document).where(Document.status == "draft",
                                                           Document.counterparty_id.is_not(None))):
                self._drafts.setdefault(d.counterparty_id, {})[d.id] = d
        return self._drafts

    def _open_reviews(self) -> list[ReviewItem]:
        if not hasattr(self, "_reviews"):
            self._reviews = list(self.s.scalars(select(ReviewItem).where(
                ReviewItem.kind.in_(("unclassified", "new_counterparty")), ReviewItem.status == "open")))
        return self._reviews

    def _find(self, tax_id: str | None, masked: str | None, name: str) -> Counterparty | None:
        idx = self._index()
        if tax_id:
            return idx["tax"].get(tax_id)
        if masked:
            return idx["masked"].get((masked, name or None)) or idx["masked"].get((masked, None))
        return idx["name"].get(name)

    def _rename(self, cp: Counterparty, new_name: str) -> None:
        old = cp.name
        cp.name = new_name
        cp.aliases = [a for a in {*(cp.aliases or []), old} if a != new_name]

    def _apply_to_drafts(self, cp: Counterparty, actor: str) -> int:
        """這個廠商等待分類的草稿：套用類別並入帳，對應的待確認一併結案。"""
        if not cp.default_category_code:
            return 0
        docs = self._drafts_by_vendor().get(cp.id, {})
        if not docs:
            return 0
        from app.feature.posting import PostingService
        posting = PostingService(self.s)
        now = datetime.now(timezone.utc)
        done = 0
        for it in [i for i in self._open_reviews() if i.kind == "unclassified" and i.status == "open"]:
            doc_id = (it.payload or {}).get("document_id")
            doc = docs.get(uuid.UUID(doc_id)) if doc_id else None
            if doc is None or doc.status != "draft":
                continue
            for ln in doc.lines:
                ln.category_code = cp.default_category_code
            posting.post(doc)
            it.status, it.resolved_by, it.resolved_at = "resolved", actor, now
            it.resolution = {"action": "set_category", "category_code": cp.default_category_code, "via": "vendor_map",
                             "document_id": str(doc.id)}
            done += 1
        return done

    def _close_new_vendor_review(self, cp: Counterparty, actor: str) -> None:
        for it in [i for i in self._open_reviews() if i.kind == "new_counterparty" and i.status == "open"
                   and i.ref == f"counterparty:{cp.id}"]:
            it.status, it.resolved_by = "resolved", actor
            it.resolution = {"action": "vendor_map", "category_code": cp.default_category_code}
            it.resolved_at = datetime.now(timezone.utc)
